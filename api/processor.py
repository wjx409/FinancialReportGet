"""处理引擎：复用 financial_report_fetcher.py 的核心查询逻辑"""
import re
import time
import uuid
import shutil
import os
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict, Any

import requests
import openpyxl
from openpyxl.styles import Font

# -------- 从 financial_report_fetcher.py 复制的常量 --------
QUERY_URL = "http://www.cninfo.com.cn/new/hisAnnouncement/query"
FILE_URL_PREFIX = "http://static.cninfo.com.cn/"
REQUEST_DELAY = 0.8
MAX_RETRY = 3
TIME_OUT = 15
TARGET_YEAR = "2025"
CATEGORY = "category_ndbg_szsh;"  # 仅年报

DEFAULT_HEADERS = {
    "Accept": "application/json, text/javascript, */*; q=0.9",
    "Accept-Language": "zh-CN,zh;q=0.9",
    "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
    "Origin": "http://www.cninfo.com.cn",
    "Referer": "http://www.cninfo.com.cn/new/disclosure",
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "X-Requested-With": "XMLHttpRequest",
}


# -------- 工具函数 --------
def normalize_stock_code(code: Any) -> str:
    if code is None:
        return ""
    raw = str(code).strip()
    if not raw:
        return ""
    m = re.search(r"\d+", raw)
    if not m:
        return ""
    return m.group(0).zfill(6)


def detect_market(stock_code: str) -> str:
    if not stock_code or len(stock_code) < 6:
        return ""
    head = stock_code[:3]
    if head in {"600", "601", "603", "605", "688", "900"}:
        return "sse"
    if head in {"000", "001", "002", "003", "300", "301", "200"}:
        return "szse"
    if stock_code[0] in {"8", "9"}:
        return "bse"
    return "szse"


def detect_plate(stock_code: str) -> str:
    market = detect_market(stock_code)
    if market == "sse":
        return "sh"
    if market == "szse":
        return "sz"
    if market == "bse":
        return "bj"
    return ""


def strip_html(text: str) -> str:
    if not text:
        return ""
    return re.sub(r"<[^>]+>", "", str(text)).strip()


def parse_timestamp(ts: Optional[int]) -> str:
    if not ts:
        return ""
    try:
        return datetime.fromtimestamp(int(ts) / 1000).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""


# -------- 核心查询函数 --------
def fetch_annual_report(code: str, company_name: Optional[str] = None) -> Dict[str, str]:
    result = {"title": "", "url": "", "publish_date": "", "status": ""}
    if not code and not (company_name and str(company_name).strip()):
        result["status"] = "缺少证券代码/公司名称"
        return result

    market = detect_market(code)
    plate = detect_plate(code)
    column = market or ""
    keyword = code if code else (company_name or "").strip()

    payload = {
        "pageNum": "1", "pageSize": "30", "column": column, "tabName": "fulltext",
        "plate": plate, "stock": "", "searchkey": keyword, "secid": "",
        "category": CATEGORY, "trade": "", "seDate": "",
        "sortName": "time", "sortType": "desc", "isHLtitle": "true",
    }

    session = requests.Session()
    last_err = ""

    for attempt in range(1, MAX_RETRY + 1):
        try:
            resp = session.post(QUERY_URL, headers=DEFAULT_HEADERS, data=payload, timeout=TIME_OUT)
            if resp.status_code != 200:
                last_err = f"HTTP {resp.status_code}"
                time.sleep(REQUEST_DELAY * attempt)
                continue

            data = resp.json()
            announcements = data.get("announcements", []) or []
            if not announcements:
                result["status"] = "未查询到定期报告"
                return result

            pre_filtered = []
            for a in announcements:
                sec_code = (a.get("secCode") or "").strip()
                sec_name = strip_html(a.get("secName") or "")
                title = strip_html(a.get("announcementTitle") or "")

                body_ok = False
                if code and sec_code == code:
                    body_ok = True
                elif (not code) and company_name:
                    name = (company_name or "").strip()
                    if name and (name in sec_name or name in title):
                        body_ok = True
                if not body_ok:
                    continue
                if TARGET_YEAR not in title:
                    continue
                if "年度报告" not in title:
                    continue
                pre_filtered.append((a, title))

            # 优先完整年报，排除摘要/英文版
            preferred, fallback = [], []
            for item in pre_filtered:
                _, t = item
                lowered = t.lower()
                bad_words = ("摘要", "英文版", "海外版", "h股", "b股",
                             "english", "英文", "外文")
                (preferred if not any(b in lowered for b in bad_words) else fallback).append(item)

            final_list = preferred if preferred else fallback
            if not final_list:
                result["status"] = f"未查询到 {TARGET_YEAR} 年年报"
                return result

            first, clean_title = final_list[0]
            adjunct = first.get("adjunctUrl", "") or ""
            ts = first.get("announcementTime", 0)
            result["title"] = clean_title
            result["url"] = FILE_URL_PREFIX + adjunct if adjunct else ""
            result["publish_date"] = parse_timestamp(ts)
            result["status"] = "成功"
            return result

        except Exception as e:
            last_err = f"异常: {e}"
            time.sleep(REQUEST_DELAY * attempt)

    result["status"] = f"查询失败: {last_err}"
    return result


# -------- Excel 读取 --------
def find_col_indexes(headers: List[Any]) -> Dict[str, int]:
    mapping = {"code": -1, "name": -1}
    for idx, h in enumerate(headers):
        if h is None:
            continue
        text = str(h).strip()
        if mapping["code"] < 0 and (
            text in {"证券代码", "股票代码", "代码", "scode", "stock_code", "stock code"}
            or re.search(r"证券\s*代码|股票\s*代码", text)
        ):
            mapping["code"] = idx
        if mapping["name"] < 0 and (
            text in {"公司名称", "机构名称", "证券简称", "简称", "公司简称", "name", "company"}
            or re.search(r"公司\s*名称|机构\s*名称", text)
        ):
            mapping["name"] = idx
    return mapping


# -------- 任务目录 --------
TEMP_DIR = Path("/tmp/finreport_tasks")
TEMP_DIR.mkdir(exist_ok=True)


def get_task_dir(task_id: str) -> Path:
    d = TEMP_DIR / task_id
    d.mkdir(exist_ok=True)
    return d


# -------- 进度回调类型 --------
ProgressCallback = callable  # def(current, total, RowResult)


def process_excel(
    input_path: str,
    task_id: str,
    on_progress: Optional[ProgressCallback] = None,
) -> tuple[List[Dict], str]:
    """
    处理 Excel 文件，返回 (结果列表, 输出文件路径)。
    on_progress(row_index, total, row_result) 每行完成后调用。
    """
    wb = openpyxl.load_workbook(input_path)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 文件为空")

    headers = list(rows[0])
    col_map = find_col_indexes(headers)

    # 追加列
    new_headers = headers[:]
    new_headers.extend(["最新财报标题", "最新财报链接", "最新财报发布时间"])
    title_col = len(headers)
    url_col = len(headers) + 1
    date_col = len(headers) + 2

    for c_idx, val in enumerate(new_headers, start=1):
        sheet.cell(row=1, column=c_idx, value=val)

    results = []
    total = len(rows) - 1

    for r_idx, row in enumerate(rows[1:], start=2):
        values = list(row) + [None] * (len(new_headers) - len(row))
        code_val = values[col_map["code"]] if col_map["code"] >= 0 else ""
        name_val = values[col_map["name"]] if col_map["name"] >= 0 else ""

        code_str = normalize_stock_code(code_val)
        company_str = str(name_val).strip() if name_val else ""

        info = fetch_annual_report(code_str, company_str)

        sheet.cell(row=r_idx, column=title_col + 1, value=info["title"])
        cell_url = sheet.cell(row=r_idx, column=url_col + 1, value=info["url"])
        if info["url"]:
            cell_url.hyperlink = info["url"]
            cell_url.font = Font(color="0563C1", underline="single")
        sheet.cell(row=r_idx, column=date_col + 1, value=info["publish_date"])

        row_result = {
            "row_index": r_idx - 1,
            "code": code_str,
            "name": company_str,
            "title": info["title"],
            "url": info["url"],
            "publish_date": info["publish_date"],
            "status": info["status"],
        }
        results.append(row_result)

        if on_progress:
            on_progress(r_idx - 1, total, row_result)

        time.sleep(REQUEST_DELAY)

    task_dir = get_task_dir(task_id)
    # 在原文件名基础上加 "_财报链接_2025"
    stem = Path(input_path).stem
    out_path = str(task_dir / f"{stem}_财报链接_{TARGET_YEAR}.xlsx")
    wb.save(out_path)
    return results, out_path
